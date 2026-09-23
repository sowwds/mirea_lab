#!/usr/bin/env python3
"""Собирает учебный отчёт: Markdown → HTML + DOCX + PDF.

PDF печатается Chromium из HTML. DOCX создаётся из cmd/report_template.docx,
поэтому титульный лист остаётся исходным шаблоном кафедры.
"""

from __future__ import annotations

import argparse
import base64
import html
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "cmd" / "report_template.docx"
CHROMIUM = shutil.which("chromium")
PDFTOTEXT = shutil.which("pdftotext")


def metadata(markdown: str) -> tuple[str, str]:
    title_match = re.search(r"^#\s+(.+)$", markdown, re.MULTILINE)
    topic_match = re.search(r"^\*\*Тема:\*\*\s*(.+)$", markdown, re.MULTILINE)
    if not title_match or not topic_match:
        raise ValueError("В начале Markdown нужны '# Практическая работа №…' и строка '**Тема:** …'.")
    return title_match.group(1).strip(), topic_match.group(1).strip()


def inline(text: str) -> str:
    text = html.escape(text)
    text = re.sub(r"`([^`]+)`", r"<code>\1</code>", text)
    return re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", text)


def parse(markdown: str) -> list[tuple[str, object]]:
    """Небольшой предсказуемый Markdown-парсер для отчётов."""
    lines = markdown.splitlines()
    blocks: list[tuple[str, object]] = []
    index = 0
    skipped_header = False
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("# ") and not skipped_header:
            skipped_header = True
            index += 1
            continue
        if line.startswith("**Тема:**"):
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            if index == len(lines):
                raise ValueError("Не закрыт блок кода ``` в Markdown.")
            blocks.append(("code", (language, code)))
            index += 1
            continue
        match = re.match(r"^(#{2,4})\s+(.+)$", line)
        if match:
            blocks.append(("heading", (len(match.group(1)) - 1, match.group(2).strip())))
            index += 1
            continue
        image = re.match(r"^!\[(.*?)\]\((.*?)\)$", line)
        if image:
            blocks.append(("image", (image.group(1).strip(), image.group(2).strip())))
            index += 1
            continue
        if line.startswith("- "):
            items: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("- "):
                items.append(lines[index].strip()[2:])
                index += 1
            blocks.append(("ul", items))
            continue
        if re.match(r"^\d+\.\s+", line):
            items = []
            while index < len(lines) and re.match(r"^\d+\.\s+", lines[index].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[index].strip()))
                index += 1
            blocks.append(("ol", items))
            continue
        paragraph = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "![", "- ")) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph.append(candidate)
            index += 1
        blocks.append(("p", " ".join(paragraph)))
    return blocks


def headings(blocks: list[tuple[str, object]]) -> list[str]:
    return [value[1] for kind, value in blocks if kind == "heading" and value[0] == 1 and value[1] != "Содержание"]


def toc_html(items: list[str], pages: dict[str, int]) -> str:
    rows = []
    for item in items:
        number = str(pages[item]) if item in pages else ""
        rows.append(f'<div class="toc-line"><span>{inline(item)}</span><span>{number}</span></div>')
    return "\n".join(rows)


def template_emblem_data_url() -> str:
    """Возвращает эмблему, встроенную в титульный шаблон, для HTML/PDF."""
    with zipfile.ZipFile(TEMPLATE) as archive:
        media = sorted(name for name in archive.namelist() if name.startswith("word/media/"))
        if not media:
            raise RuntimeError("В cmd/report_template.docx не найдена эмблема в word/media.")
        filename = media[0]
        extension = Path(filename).suffix.lower().lstrip(".")
        mime = "image/jpeg" if extension in {"jpg", "jpeg"} else f"image/{extension}"
        encoded = base64.b64encode(archive.read(filename)).decode("ascii")
    return f"data:{mime};base64,{encoded}"


def render_html(title: str, topic: str, blocks: list[tuple[str, object]], pages: dict[str, int]) -> str:
    figure_no = 0
    listing_no = 0
    section = "0"
    output: list[str] = []
    top_headings = headings(blocks)
    emblem = template_emblem_data_url()
    for kind, value in blocks:
        if kind == "heading":
            level, text = value
            if level == 1:
                if text == "Содержание":
                    output.append('<h1 id="toc">СОДЕРЖАНИЕ</h1>')
                    output.append('<div class="toc">' + toc_html(top_headings, pages) + "</div>")
                else:
                    match = re.match(r"^(\d+)\s+", text)
                    if match:
                        section = match.group(1)
                        figure_no = 0
                        listing_no = 0
                    output.append(f"<h1>{inline(text)}</h1>")
            else:
                output.append(f"<h{level + 1}>{inline(text)}</h{level + 1}>")
        elif kind == "p":
            output.append(f"<p>{inline(value)}</p>")
        elif kind in {"ul", "ol"}:
            output.append(f"<{kind}>" + "".join(f"<li>{inline(item)}</li>" for item in value) + f"</{kind}>")
        elif kind == "image":
            caption, source = value
            figure_no += 1
            figure_id = f"{section}.{figure_no}"
            output.append(
                '<figure><img src="' + html.escape(source) + '" alt="' + html.escape(caption) + '">'
                + f"<figcaption>Рисунок {figure_id} — {inline(caption)}</figcaption></figure>"
            )
        elif kind == "code":
            language, code = value
            if language.lower() in {"text", "plaintext"}:
                output.append(f'<div class="formula">{html.escape(chr(10).join(code))}</div>')
                continue
            listing_no += 1
            output.append(
                f'<div class="listing-caption">Листинг {section}.{listing_no} — Фрагмент программы'
                f' на {html.escape(language or "C")}</div><pre><code>{html.escape(chr(10).join(code))}</code></pre>'
            )
    return f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>
@page {{ size: A4; margin: 20mm 10mm 20mm 30mm; }}
body {{ font: 14pt/1.5 'Times New Roman', serif; color: #000; }}
.title {{ break-after: page; height: 247mm; display: flex; flex-direction: column; text-align: center; }}
.title .institution {{ margin-top: 12mm; }} .title .emblem {{ height: 29mm; width: auto; margin: 5mm auto 0; }} .title .work {{ margin-top: auto; }}
.title .author {{ margin: auto 0 0 52%; text-align: left; }} .title .city {{ margin-top: auto; }}
h1 {{ font-size: 18pt; text-transform: uppercase; text-align: center; break-before: page; break-after: avoid; margin: 0 0 10mm; }}
#toc {{ break-before: auto; }} h2 {{ font-size: 16pt; margin: 15mm 0 10mm; break-after: avoid; }}
h3 {{ font-size: 14pt; margin: 12mm 0 6mm; break-after: avoid; }}
p {{ margin: 0; text-align: justify; text-indent: 1.25cm; }} li {{ text-align: justify; overflow-wrap: anywhere; word-break: break-all; }}
.toc-line {{ display: grid; grid-template-columns: 1fr 12mm; gap: 4mm; }} .toc-line span:first-child {{ border-bottom: 1px dotted #000; }}
figure {{ break-inside: avoid; margin: 8mm auto; text-align: center; }} figure img {{ max-height: 170mm; max-width: 100%; }}
figcaption {{ font-size: 12pt; font-weight: bold; text-align: center; }}
pre {{ border: 1px solid #000; padding: 3mm; white-space: pre-wrap; font: 10pt/1.1 'Courier New', monospace; }}
.listing-caption {{ margin-top: 6mm; font-size: 12pt; font-style: italic; break-after: avoid; }}
.formula {{ margin: 5mm 0; white-space: pre-wrap; text-align: center; font-family: 'Times New Roman', serif; }}
</style></head><body>
<section class="title"><div class="institution">МИНОБРНАУКИ РОССИИ<br>Федеральное государственное бюджетное образовательное учреждение высшего образования<br>«МИРЭА — Российский технологический университет»<br>РТУ МИРЭА</div><img class="emblem" src="{emblem}" alt="Эмблема РТУ МИРЭА"><div class="work">ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ<br>«{html.escape(topic)}»<br>по дисциплине «Программирование киберфизических систем»</div><div class="author">Выполнил студент группы ЭФБО-10-23<br>Ефремов А.И.<br><br>Принял к.т.н., доцент<br>Сухатерин А.Б.</div><div class="city">Москва 2026</div></section>
{''.join(output)}</body></html>"""


def set_font(run, size: int = 14, bold: bool | None = None, name: str = "Times New Roman") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def setup_document(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(14)


def all_paragraphs(document: Document):
    for paragraph in document.paragraphs:
        yield paragraph
    def table_paragraphs(table):
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from table_paragraphs(nested)
    for table in document.tables:
        yield from table_paragraphs(table)


def replace_template_title(document: Document, topic: str) -> None:
    for paragraph in all_paragraphs(document):
        if "Название практической" in paragraph.text:
            paragraph.text = paragraph.text.replace("Название практической", topic)
            for run in paragraph.runs:
                set_font(run, size=14)


def set_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Cm(1.25) if first_line else None


def add_docx_text(document: Document, text: str, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True) -> None:
    paragraph = document.add_paragraph()
    set_paragraph(paragraph, align, first_line)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        run = paragraph.add_run(part.strip("*"))
        set_font(run, bold=part.startswith("**"))


def add_docx_heading(document: Document, level: int, text: str, page_break: bool) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    # Не записываем w:pageBreakBefore с w:val="0": OnlyOffice может
    # воспринять такой элемент как разрыв страницы.
    if page_break:
        fmt.page_break_before = True
    fmt.keep_with_next = True
    fmt.space_before = Pt(0 if level == 1 else 24)
    fmt.space_after = Pt(12)
    fmt.line_spacing = 1.5
    run = paragraph.add_run(text.upper() if level == 1 else text)
    set_font(run, size=18 if level == 1 else 16 if level == 2 else 14, bold=True)


def add_docx_image(document: Document, source: Path, caption: str, number: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(source), width=Cm(14.5))
    caption_p = document.add_paragraph()
    set_paragraph(caption_p, WD_ALIGN_PARAGRAPH.CENTER, False)
    run = caption_p.add_run(f"Рисунок {number} — {caption}")
    set_font(run, size=12, bold=True)


def add_docx_listing(document: Document, language: str, code: list[str], number: str) -> None:
    caption = document.add_paragraph()
    set_paragraph(caption, WD_ALIGN_PARAGRAPH.LEFT, False)
    run = caption.add_run(f"Листинг {number} — Фрагмент программы на {language or 'C'}")
    set_font(run, size=12)
    run.italic = True
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    paragraph = table.cell(0, 0).paragraphs[0]
    set_paragraph(paragraph, WD_ALIGN_PARAGRAPH.LEFT, False)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("\n".join(code))
    set_font(run, size=10, name="Courier New")


def add_docx_formula(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    set_paragraph(paragraph, WD_ALIGN_PARAGRAPH.CENTER, False)
    run = paragraph.add_run("\n".join(lines))
    set_font(run)


def render_docx(markdown_path: Path, title: str, topic: str, blocks, pages: dict[str, int], output: Path) -> None:
    document = Document(str(TEMPLATE))
    setup_document(document)
    replace_template_title(document, topic)
    # Не добавляем пустой абзац или новую секцию после титульного листа: при
    # почти заполненном шаблоне OnlyOffice переносил такой абзац на отдельный
    # пустой лист. Первый заголовок сам создаёт единственный разрыв страницы.
    top_headings = headings(blocks)
    section = "0"
    figure_no = 0
    listing_no = 0
    for kind, value in blocks:
        if kind == "heading":
            level, text = value
            if level == 1:
                add_docx_heading(document, 1, text, page_break=True)
                match = re.match(r"^(\d+)\s+", text)
                if match:
                    section = match.group(1)
                    figure_no = 0
                    listing_no = 0
                if text == "Содержание":
                    for item in top_headings:
                        paragraph = document.add_paragraph()
                        set_paragraph(paragraph, WD_ALIGN_PARAGRAPH.LEFT, False)
                        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(16), 2, 1)
                        left = paragraph.add_run(item)
                        set_font(left)
                        paragraph.add_run("\t")
                        right = paragraph.add_run(str(pages.get(item, "")))
                        set_font(right)
            else:
                add_docx_heading(document, level, text, page_break=False)
        elif kind == "p":
            add_docx_text(document, value)
        elif kind in {"ul", "ol"}:
            for position, item in enumerate(value, start=1):
                add_docx_text(document, ("– " if kind == "ul" else f"{position}. ") + item, first_line=False)
        elif kind == "image":
            caption, source = value
            figure_no += 1
            add_docx_image(document, markdown_path.parent / source, caption, f"{section}.{figure_no}")
        elif kind == "code":
            language, code = value
            if language.lower() in {"text", "plaintext"}:
                add_docx_formula(document, code)
                continue
            listing_no += 1
            add_docx_listing(document, language, code, f"{section}.{listing_no}")
    document.save(output)


def chromium_pdf(html_path: Path, pdf_path: Path) -> None:
    if not CHROMIUM:
        raise RuntimeError("Chromium не найден: невозможно сформировать PDF.")
    with tempfile.TemporaryDirectory(prefix="report-chromium-") as profile:
        command = [CHROMIUM, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer", f"--user-data-dir={profile}", f"--print-to-pdf={pdf_path}", html_path.as_uri()]
        subprocess.run(command, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)


def pages_from_pdf(pdf_path: Path, expected: list[str]) -> dict[str, int]:
    if not PDFTOTEXT:
        return {}
    extracted = subprocess.run([PDFTOTEXT, "-layout", str(pdf_path), "-"], check=True, text=True, stdout=subprocess.PIPE).stdout
    result: dict[str, int] = {}
    for number, page in enumerate(extracted.split("\f"), start=1):
        compact = re.sub(r"\s+", " ", page).strip().upper()
        # На странице оглавления встречаются все заголовки. Она не является
        # началом ни одного из разделов, поэтому не учитывается.
        if "СОДЕРЖАНИЕ" in compact:
            continue
        for heading in expected:
            if heading.upper() in compact and heading not in result:
                result[heading] = number
    return result


def build(markdown_path: Path) -> None:
    markdown = markdown_path.read_text(encoding="utf-8")
    title, topic = metadata(markdown)
    blocks = parse(markdown)
    output_base = markdown_path.with_suffix("")
    html_path = output_base.with_suffix(".html")
    docx_path = output_base.with_suffix(".docx")
    pdf_path = output_base.with_suffix(".pdf")
    expected = headings(blocks)
    html_path.write_text(render_html(title, topic, blocks, {}), encoding="utf-8")
    chromium_pdf(html_path, pdf_path)
    pages = pages_from_pdf(pdf_path, expected)
    html_path.write_text(render_html(title, topic, blocks, pages), encoding="utf-8")
    chromium_pdf(html_path, pdf_path)
    render_docx(markdown_path, title, topic, blocks, pages, docx_path)
    print(f"HTML: {html_path}")
    print(f"DOCX: {docx_path}")
    print(f"PDF:  {pdf_path}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Собрать отчёт из Markdown.")
    parser.add_argument("markdown", type=Path, help="например prac1/report_prac1.md")
    args = parser.parse_args()
    build(args.markdown.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
